"""Adapter for a PTAT xlsx/CSV export -- one row per ~100ms sample, ~800 columns wide (see
gitignorethis/PTAT-Visualization/README.md for the column-naming conventions this relies on).

Recognized by the presence of a "Relative Time(mS)" column, which is also what's used to
reconstruct absolute per-row timestamps: only the FIRST row's Date+Time is parsed (PTAT's own
"HH:MM:SS:mmm" time format, colon-separated millis rather than the usual dot), and every other
row's timestamp is that origin plus its own Relative Time(mS) offset -- cheaper and exactly as
accurate as parsing all ~2000 rows individually, since Relative Time(mS) is already a precise
elapsed-ms counter from the same clock.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Optional

import pandas as pd

from .base import Trace, split_static

REL_TIME_COL = "Relative Time(mS)"
_DROP_COLS = (REL_TIME_COL, "Date", "Time", "Diff time(mS)")


def sniff(path: Path) -> bool:
    if path.suffix.lower() not in (".xlsx", ".csv"):
        return False
    try:
        header = _read_header(path)
    except Exception:
        return False
    return REL_TIME_COL in header


def _read_header(path: Path) -> list:
    if path.suffix.lower() == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        wb.close()
        return [str(c) for c in row]
    return pd.read_csv(path, nrows=0).columns.tolist()


def _parse_ptat_datetime(date_value: str, time_value: str) -> datetime:
    day, month, year = (int(x) for x in str(date_value).split("/"))
    hour, minute, second, millis = (int(x) for x in str(time_value).split(":"))
    return datetime(year, month, day, hour, minute, second, millis * 1000)


def load(path: Path) -> Optional[Trace]:
    if path.suffix.lower() == ".xlsx":
        df = pd.read_excel(path, sheet_name=0, engine="openpyxl")
    else:
        df = pd.read_csv(path, low_memory=False)
    if REL_TIME_COL not in df.columns or df.empty:
        return None

    rel_ms = pd.to_numeric(df[REL_TIME_COL], errors="coerce")
    df = df.loc[rel_ms.notna()].copy()
    rel_ms = rel_ms.loc[rel_ms.notna()]
    origin = _parse_ptat_datetime(df["Date"].iloc[0], df["Time"].iloc[0])
    df["_ts"] = pd.Timestamp(origin) + pd.to_timedelta(rel_ms - rel_ms.iloc[0], unit="ms")
    df = df.sort_values("_ts").reset_index(drop=True)

    value_df, static_info = split_static(df, keep=("_ts",), drop=_DROP_COLS)
    return Trace(name="PTAT", kind="sampled", df=value_df, source_path=str(path), static_info=static_info)
